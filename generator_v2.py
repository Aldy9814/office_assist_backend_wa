import pandas as pd
import os
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from xhtml2pdf import pisa
import sqlite3

def get_db_connection():
    return sqlite3.connect('office_data.db')

def handle_tarik_absen(conn, keyword):
    query = f"SELECT * FROM attendance WHERE tanggal LIKE '{keyword}-%'"
    print(f"Menjalankan query: {query}")
    df = pd.read_sql_query(query, conn)

    filename_base = f"Laporan_Absen_{keyword}" if keyword and keyword.lower() != "null" else "Laporan_Absen"
    judul_laporan = "Laporan Kehadiran Karyawan"

    return df, filename_base, judul_laporan

def handle_tarik_inventaris(conn, keyword):
    if keyword and keyword.lower() != "null":
        search_term = f"%{keyword}%"
        query = f"SELECT * FROM inventaris WHERE nama_barang LIKE '{search_term}' OR kategori LIKE '{search_term}' OR pemegang LIKE '{search_term}'"
        print(f"Menjalankan query: {query}")
        df = pd.read_sql_query(query, conn)
        filename_base = f"Laporan_Inventaris_{keyword}"
    else:
        query = "SELECT * FROM inventaris"
        print(f"Menjalankan query: {query}")
        df = pd.read_sql_query(query, conn)
        filename_base = "Laporan_Inventaris_Lengkap"

    judul_laporan = "Laporan Inventaris Aset"

    return df, filename_base, judul_laporan


INTENT_ROUTES = {
    "tarik_absen": handle_tarik_absen,
    "tarik_inventaris": handle_tarik_inventaris
}

def generate_report(intent, file_format, keyword=None):
    handler_function = INTENT_ROUTES.get(intent)

    if not handler_function:
        print(f"Error: Fitur untuk intent '{intent}' belum tersedia.")
        return None

    conn = get_db_connection()
    try:
        df, filename_base, judul_laporan = handler_function(conn, keyword)
    finally:
        conn.close()

    print(f"JUMLAH DATA DITEMUKAN: {len(df)}")

    if df.empty:
        print("Data kosong, tidak ada file yang di-generate.")
        return None

    filter_val = keyword

    try:
        if file_format.lower() == "excel":
            file_path = f"{filename_base}.xlsx"
            writer = pd.ExcelWriter(file_path, engine='openpyxl')
            df.to_excel(writer, index=False, sheet_name='Laporan')

            workbook = writer.book
            worksheet = writer.sheets['Laporan']

            header_fill = PatternFill(start_color="203764", end_color="203764", fill_type="solid") 
            header_font = Font(color="FFFFFF", bold=True)
            border = Border(
                left=Side(style='thin', color='D4D4D4'),
                right=Side(style='thin', color='D4D4D4'),
                top=Side(style='thin', color='D4D4D4'),
                bottom=Side(style='thin', color='D4D4D4')
            )

            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.border = border

            for i, col_name in enumerate(df.columns, start=1):
                col_letter = get_column_letter(i)
                if "Nama" in str(col_name) or "Kategori" in str(col_name):
                    worksheet.column_dimensions[col_letter].width = 25
                else:
                    worksheet.column_dimensions[col_letter].width = 15

            for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row, min_col=1, max_col=worksheet.max_column):
                for cell in row:
                    cell.border=border
                    cell.alignment = Alignment(horizontal='center', vertical='center')

            writer.close()
            print(f"Success! Excel is saved in: {os.path.abspath(file_path)}")
            return file_path

        elif file_format.lower() == "pdf":
            file_path = f"{filename_base}.pdf"
            html_table = df.to_html(index=False, border=1, justify="center")
            periode = filter_val if filter_val and filter_val != "null" else "Semua Data"

            html_content = f"""
            <html>
            <head>
                <style>
                    @page {{ size: A4 landscape; margin: 1cm; }}
                    body {{ font-family: Helvetica, Arial, sans-serif; font-size: 10pt; }}
                    .header-title {{ text-align: center; color: #203764; font-size: 24pt; font-weight: bold; margin-bottom: 5px; }}
                    .header-subtitle {{ text-align: center; color: #555555; font-size: 12pt; margin-bottom: 20px; border-bottom: 2px solid #203764; padding-bottom: 10px; }}
                    table {{ width: 100%; border: 0.5px solid #000000; }}
                    th {{ background-color: #203764; color: #ffffff; padding: 6px; text-align: center; font-weight: bold; border: 0.5px solid #000000; }}
                    td {{ padding: 5px; border: 0.5px solid #000000; text-align: center; }}
                </style>
            </head>
            <body>
                <div class="header-title">{judul_laporan}</div>
                <div class="header-subtitle">Periode/Filter: {periode}</div>
                {html_table}
            </body>
            </html>
            """

            with open(file_path, "w+b") as result_file:
                pisa_status = pisa.CreatePDF(html_content, dest=result_file)

            if pisa_status.err:
                print("Error when converting HTML to PDF")
                return None
            else:
                print(f"Success! PDF is saved in: {os.path.abspath(file_path)}")
                return file_path
                
    except Exception as e:
        print(f"Error saat meng-generate laporan: {e}")
        return None

    